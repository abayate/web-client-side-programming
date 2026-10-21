import json
from docx import Document
from docx.shared import Pt, Inches

U = json.load(open('urls.json'))
NAME = 'Anthony Bayate'


def base(title, sub):
    d = Document()
    st = d.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    d.add_heading(title, 0)
    d.add_paragraph(f'{NAME}  |  Kean University  |  Web Client-Side Programming  |  {sub}')
    return d


def bullets(d, items):
    for i in items:
        d.add_paragraph(i, style='List Bullet')


def table(d, headers, rows):
    t = d.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for r in rows:
        c = t.add_row().cells
        for i, v in enumerate(r):
            c[i].text = v


# ---------- 1. Create a Game in JS: links only ----------
d = base('Create a Game in JS: Live Links', 'Games 1-4 (Tasks 1-7)')
table(d, ['Game', 'Tasks completed', 'Link'], [
    ('Game 1: Nebula Defender (space shooter)',
     'Task 1 (rapid-fire power-up, top-5 localStorage leaderboard); Task 2 (boss at 1000 pts, custom look, name on GUI)', U['game1']),
    ('Game 2: Synthwave Breakout',
     'Task 3 (P to pause, falling power-ups: wide / slow / sticky / extra life); Task 4 (top-5 initials leaderboard, custom look, name on GUI)', U['game2']),
    ('Game 3: Labyrinth Escape',
     'Task 5 (coins gate the exit, patrolling guards); Task 6 (time-boost, freeze and speed power-ups, custom look, name on GUI)', U['game3']),
    ('Game 4: Void Runner (my own game)',
     'Task 7 (double-jump, dash, combo, near-miss bonus, magnet and slow-time power-ups, custom look, name on GUI)', U['game4']),
])
d.add_paragraph()
d.add_paragraph('Per the instructions, only links are submitted (no source code).')
d.save('Create_a_Game_in_JS_Links.docx')

# ---------- 2. Week 5 AS2 ----------
d = base('Week 5 AS2: Class Activity and Reflection', 'Neon Defense (tower defense)')
d.add_heading('Live game', 1)
d.add_paragraph(U['as2'])

d.add_heading('What I built', 1)
d.add_paragraph(
    'Neon Defense is a tower-defense game written from scratch in plain HTML, CSS and JavaScript '
    '(index.html, style.css, game.js, no libraries). Enemies follow a winding neon path to the core. '
    'The player spends credits on four tower types (Pulse, Cannon with splash damage, Frost that slows, '
    'Sniper with long range), can upgrade each tower to level 3 or sell it, and has to survive 20 waves. '
    'Every fifth wave ends with a boss. The best wave reached is saved with localStorage. I avoided simple '
    'games on purpose: this one has pathing, targeting, wave scaling, an economy and upgrade/sell decisions.')
d.add_heading('How I used HTML, CSS and JS', 1)
bullets(d, [
    'HTML: page structure, the shop, stats panel and start/game-over overlay.',
    'CSS: dark neon theme, grid layout for the shop and stats, hover/selected states, responsive canvas.',
    'JS: canvas rendering, a requestAnimationFrame game loop, enemy path-following, tower targeting '
    '(furthest enemy along the path), projectiles, splash and slow effects, a wave generator with HP scaling, '
    'DOM-based UI and keyboard shortcuts (1-4, Space, P, Esc).',
])
d.add_heading('How I built it with an LLM (step by step)', 1)
bullets(d, [
    'Asked the LLM to walk me through the project: pick a non-trivial game, plan the three files, and describe the game loop before writing code.',
    'Built it in stages: map + path, enemies moving, towers shooting, then economy, waves, upgrades, polish.',
    'Ran it after each stage and fed problems back. Example: late waves were too hard, so I lowered the HP growth from 1.17x to 1.14x per wave.',
    'Ran a scripted playthrough in the browser console (places towers, fast-forwards waves) to check the whole loop for errors before deploying.',
    'Deployed with the Vercel CLI: install it, run "vercel login", then "vercel --prod" in the project folder. Vercel serves the static files directly, so no build step is needed.',
])
d.add_heading('Peer testing', 1)
d.add_paragraph('(Fill in after classmates play. Suggested table below.)')
table(d, ['Tester', 'What they said / found', 'What I changed'], [('', '', ''), ('', '', ''), ('', '', '')])
d.add_heading('Reflection', 1)
d.add_paragraph('(Write 1-2 short paragraphs in your own words. Starting points: what was hardest about the game loop or '
                'tower targeting; what the LLM got right or wrong; what you would add next, such as more maps or a flying enemy.)')
d.save('Week5_AS2_Class_Activity_and_Reflection.docx')

# ---------- 3. Lab 6 ----------
d = base('Week 6 Lab #6: Play with AI', 'Fleet Command: Battleship vs. Admiral Orca')
d.add_heading('Live game', 1)
d.add_paragraph(U['as1'])

d.add_paragraph('Team: ' + NAME + ' (individual submission)')
d.add_heading('Idea', 1)
d.add_paragraph(
    'Fleet Command is Battleship where the opponent is an AI admiral, "Admiral Orca". The player places a fleet of '
    'five ships (manually with a rotate preview, or randomly), then trades shots with the AI on 10x10 grids. '
    'Extras: a Sonar ping (3 per game) that reports how many enemy ship squares are in a 3x3 area at the cost of a turn, '
    'a selectable AI difficulty, a battle log, and a win/loss record per difficulty saved in localStorage.')
d.add_heading('The AI opponent', 1)
d.add_paragraph('The AI lives in ai.js and has three levels:')
bullets(d, [
    'Easy: fires at a random unexplored square.',
    'Medium: hunt/target. Searches on a checkerboard (every ship is at least 2 long, so half the board is enough); after a hit it probes neighbours and prefers extending a line of hits.',
    'Hard: probability-density search. For every unexplored square it counts how many legal placements of the ships still afloat would cover it and fires at the highest count. Placements that explain unresolved hits get much more weight, and sunk ships are removed from the remaining list so the model updates after every result.',
])
d.add_paragraph("The AI only sees what a human would see (hit, miss, sunk and the length of the sunk ship). It never reads the player's board.")
d.add_heading('Does the AI get smarter? Measured results', 1)
d.add_paragraph('I simulated 300 games per level against random fleets and counted the shots needed to sink all five ships (lower is better):')
table(d, ['AI level', 'Average shots to win'], [('Easy', '95.2'), ('Medium', '52.3'), ('Hard', '44.8')])
d.add_paragraph('Easy behaves like a random shooter, which needs close to 95 shots to clear the board; Hard needs under half of that.')
d.add_heading('Screenshot', 1)
d.add_picture('img/fleet.png', width=Inches(6.3))
d.add_heading('Tech and AI usage', 1)
bullets(d, [
    'HTML/CSS/JS only (index.html, style.css, ai.js, game.js). Deployed to Vercel.',
    'I used an LLM (Claude) to help design the AI levels and write code, then tested it with the shot-count simulation above and a scripted full playthrough in the browser.',
])
d.add_heading('Team / testing notes', 1)
d.add_paragraph('(Add any teammates and what classmates thought of the AI.)')
d.save('Week6_Lab6_Play_with_AI.docx')
print('ok')
